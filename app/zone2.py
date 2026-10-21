import csv
import io
import shutil
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Dict, List, Optional

from app import consent, data_dictionary, redcap, storage
from app.crud import patient_application_files as files_crud
from app.crud import patient_applications as applications_crud
from app.crud import patients as patients_crud
from app.crud import projects as projects_crud
from app.errors import NotFoundError, ValidationError
from app.logging_setup import get_logger
from app.project_schemas import Candidate, Project, ProjectRelease
from app.schemas import PatientApplicationFileUpdate

log = get_logger(__name__)

RELEASED_STATUSES = ("submitted", "approved")
ENDED = "ended"
SAMPLE_ROWS = 5
AGE_CAP = 90

PATIENT_COLUMNS = [
    "project_patient_id", "age_at_registration", "registration_date",
    "death_date", "state", "country",
]
APPLICATION_COLUMNS = [
    "project_application_id", "project_patient_id", "status",
    "created_date", "submitted_date",
]
DOCUMENT_COLUMNS = [
    "project_document_id", "project_application_id", "project_patient_id",
    "file_type", "modality", "document_date", "file_size_bytes", "created_date", "file",
]

README = """Zone 2 safe copy for project {name} ({code}), release {release}.

patients.csv      one row per patient, by project ID
applications.csv  their submitted applications
documents.csv     the documents on those applications{documents_note}
redcap_<name>.csv REDCap data, one file per dataset: coded, number and date
                  fields only; identifier and free-text fields are left out
data_dictionary.csv  what every column means

Patient IDs are project-only and cannot be matched to CHSS IDs here. Every
date of a patient is moved by the same number of days (up to 5 years either
way), so intervals between a patient's dates are true; calendar dates are
not. Date of birth is given as age at registration ({cap}+ for {cap} and
over). Names, addresses, phone numbers, emails and all free text are left out.
"""


def _submitted_applications(cursor, patient_id: str):
    return [
        a
        for a in applications_crud.list_applications(cursor, patient_id)
        if a.status in RELEASED_STATUSES
    ]


def _documents(cursor, application_id: str):
    return [
        f
        for f in files_crud.list_files(cursor, application_id)
        if f.review_status != "deleted"
        and f.deid_status == "done"
        and f.de_identified_file_path
    ]


def candidates(cursor) -> List[Candidate]:
    counts: Dict[str, List[int]] = {}
    for application in applications_crud.list_applications(cursor):
        if application.status not in RELEASED_STATUSES:
            continue
        entry = counts.setdefault(application.patient_id, [0, 0])
        entry[0] += 1
        entry[1] += len(_documents(cursor, application.id))
    return [
        Candidate(patient_id=code, applications=apps, documents=docs)
        for code, (apps, docs) in sorted(counts.items())
    ]


def _as_date(value) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def shift_date(value, days: int) -> str:
    day = _as_date(value)
    return (day + timedelta(days=days)).isoformat() if day else ""


def age_at(birth, on) -> str:
    born, then = _as_date(birth), _as_date(on)
    if not born or not then:
        return ""
    years = then.year - born.year - ((then.month, then.day) < (born.month, born.day))
    if years < 0:
        return ""
    return f"{AGE_CAP}+" if years >= AGE_CAP else str(years)


def _write_csv(path: Path, columns, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: ("" if v is None else v) for k, v in row.items()})


def _read_csv(path: Path) -> List[dict]:
    if not path.is_file():
        return []
    with open(path, newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _model_rows(models) -> List[dict]:
    return [m.model_dump(mode="json") for m in models]


def _new_release_id() -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    return f"R{stamp}-{projects_crud.random_code(4)}"


class DocumentError(ValueError):
    pass


def _clean_pdf(source: Path, target: Path, code: str, project_id: str) -> None:
    import fitz

    document = fitz.open(str(source))
    try:
        for page in document:
            hits = page.search_for(code)
            for rect in hits:
                page.add_redact_annot(rect, fill=(1, 1, 1))
            if hits:
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)
                for rect in hits:
                    page.insert_text(
                        fitz.Point(rect.x0, rect.y1 - 1),
                        project_id,
                        fontname="helv",
                        fontsize=max(6.0, (rect.y1 - rect.y0) * 0.8),
                    )
        document.set_metadata({})
        try:
            document.del_xml_metadata()
        except Exception:
            pass
        document.save(str(target), garbage=4, deflate=True)
    finally:
        document.close()

    check = fitz.open(str(target))
    try:
        leaks = any(page.search_for(code) for page in check)
        leaks = leaks or any(code in str(v) for v in (check.metadata or {}).values())
    finally:
        check.close()
    if leaks:
        raise DocumentError("the CHSS code is still in the PDF")


def _docx_paragraphs(document):
    def walk(container):
        for paragraph in container.paragraphs:
            yield paragraph
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(document)
    for section in document.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            yield from walk(part)


DOCX_PROPERTIES = (
    "author", "category", "comments", "content_status", "identifier", "keywords",
    "language", "last_modified_by", "subject", "title", "version",
)


def _clean_docx(source: Path, target: Path, code: str, project_id: str) -> None:
    import docx

    document = docx.Document(str(source))
    for paragraph in _docx_paragraphs(document):
        if code in paragraph.text and paragraph.runs:
            text = paragraph.text.replace(code, project_id)
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ""
    properties = document.core_properties
    for name in DOCX_PROPERTIES:
        try:
            setattr(properties, name, "")
        except Exception:
            pass
    document.save(str(target))

    check = docx.Document(str(target))
    text = " ".join(p.text for p in _docx_paragraphs(check))
    props = " ".join(str(getattr(check.core_properties, n, "") or "") for n in DOCX_PROPERTIES)
    if code in text or code in props:
        raise DocumentError("the CHSS code is still in the Word document")


def _shift_dicom_dates(dataset, days: int) -> None:
    for element in list(dataset.iterall()):
        if element.VR not in ("DA", "DT") or not element.value:
            continue
        values = element.value if element.VM > 1 else [element.value]
        shifted = []
        for value in values:
            text = str(value)
            day = None
            try:
                day = datetime.strptime(text[:8], "%Y%m%d").date()
            except ValueError:
                pass
            shifted.append((day + timedelta(days=days)).strftime("%Y%m%d") + text[8:] if day else "")
        element.value = shifted if element.VM > 1 else shifted[0]


def _new_uids(dataset, salt: str, depth: int = 0) -> None:
    from pydicom._uid_dict import UID_dictionary
    from pydicom.uid import generate_uid

    if depth > 8:
        return
    for element in dataset:
        if element.VR == "SQ":
            for item in element.value or []:
                _new_uids(item, salt, depth + 1)
        elif element.VR == "UI" and element.keyword not in (
            "ImplementationClassUID", "TransferSyntaxUID", "CodingSchemeUID"
        ):
            values = element.value if element.VM > 1 else [element.value]
            new = [
                str(v) if not v or str(v) in UID_dictionary
                else str(generate_uid(entropy_srcs=[salt, str(v)]))
                for v in values
            ]
            element.value = new if element.VM > 1 else new[0]


def _clean_dicom(
    source: Path, target: Path, code: str, project_id: str, days: int, uid_salt: str = ""
) -> None:
    import pydicom
    from pydicom.uid import generate_uid

    dataset = pydicom.dcmread(str(source), force=True)
    dataset.PatientID = project_id
    dataset.PatientName = project_id
    dataset.DeidentificationMethod = "CHSS de-identification; project IDs; dates shifted"
    try:
        dataset.remove_private_tags()
    except Exception:
        pass
    for element in list(dataset.iterall()):
        if isinstance(element.value, str) and code in element.value:
            element.value = element.value.replace(code, project_id)
    _shift_dicom_dates(dataset, days)
    if uid_salt:
        _new_uids(dataset, uid_salt)
        meta = getattr(dataset, "file_meta", None)
        if meta is not None and "MediaStorageSOPInstanceUID" in meta:
            meta.MediaStorageSOPInstanceUID = generate_uid(
                entropy_srcs=[uid_salt, str(meta.MediaStorageSOPInstanceUID)]
            )
    dataset.save_as(str(target))

    check = pydicom.dcmread(str(target), force=True)
    if any(isinstance(e.value, str) and code in e.value for e in check.iterall()):
        raise DocumentError("the CHSS code is still in the DICOM header")


def _copy_document(
    record, target_dir: Path, document_id: str, code: str, project_id: str, days: int,
    uid_salt: str = "",
) -> str:
    source = storage.resolve_stored_path(record.de_identified_file_path)
    if not source.is_file():
        raise DocumentError("the de-identified copy is missing from disk")
    suffix = source.suffix.lower()
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{document_id}{suffix}"
    if suffix == ".pdf":
        _clean_pdf(source, target, code, project_id)
    elif suffix in (".doc", ".docx"):
        _clean_docx(source, target, code, project_id)
    elif suffix in (".dcm", ".dicom"):
        _clean_dicom(source, target, code, project_id, days, uid_salt)
    else:
        raise DocumentError(f"'{suffix}' documents cannot be released")
    return target.name


def _refuse_rejected_documents(cursor, cohort) -> None:
    waiting = [
        f"{record.deidentified_file_name or record.original_file_name}"
        for member in cohort
        for application in _submitted_applications(cursor, member.patient_id)
        for record in _documents(cursor, application.id)
        if record.review_status == "rejected"
    ]
    if waiting:
        raise ValidationError(
            f"{len(waiting)} document(s) are waiting in Rejections; replace or "
            "delete them first: " + ", ".join(waiting[:10])
        )


def _refuse_leaded_documents(cursor, cohort) -> None:
    leaded = [
        f"{member.patient_id}: {record.deidentified_file_name or record.original_file_name}"
        for member in cohort
        for application in _submitted_applications(cursor, member.patient_id)
        for record in _documents(cursor, application.id)
        if record.leaded_phi
    ]
    if leaded:
        raise ValidationError(
            f"{len(leaded)} image(s) are marked as having leaded identifiers and "
            "can't go to Zone 2. Scrub them by hand and replace them, then clear "
            "the mark, or take their patients out of the cohort: "
            + ", ".join(leaded[:10])
        )


def _require_ready(project: Project, cohort) -> None:
    problems = []
    if project.status == ENDED:
        problems.append("the project has ended")
    if project.status == "in_review":
        problems.append("a release is already waiting for review")
    if not project.approval_reference:
        problems.append("the approval reference is missing")
    if not project.consent_confirmed:
        problems.append("consent has not been confirmed")
    if not cohort:
        problems.append("no patients are in the project")
    if problems:
        raise ValidationError("The project cannot be prepared: " + "; ".join(problems))


def prepare(cursor, project_id: str, actor_id: str) -> ProjectRelease:
    project = projects_crud.get_project_or_404(cursor, project_id)
    cohort = projects_crud.list_cohort(cursor, project.id)
    _require_ready(project, cohort)
    consent.refuse_excluded(cursor, project, [m.patient_id for m in cohort])
    if project.documents_approved:
        _refuse_rejected_documents(cursor, cohort)
        _refuse_leaded_documents(cursor, cohort)

    release_id = _new_release_id()
    archive = storage.projects_root() / project.short_code / release_id
    inbox = storage.zone2_inbox_root() / project.short_code / release_id
    safe = archive / "zone2"

    try:
        counts = _build(cursor, project, cohort, release_id, archive, safe)
        shutil.copytree(safe, inbox)
    except Exception as exc:
        shutil.rmtree(archive, ignore_errors=True)
        shutil.rmtree(inbox, ignore_errors=True)
        if isinstance(exc, DocumentError):
            raise ValidationError(f"The project cannot be prepared: {exc}") from exc
        raise

    release = projects_crud.create_release(
        cursor,
        release_id,
        project.id,
        actor_id,
        counts,
        {"archive": str(archive), "inbox": str(inbox)},
        project.documents_approved,
    )
    projects_crud.set_project_status(cursor, project.id, "in_review")
    log.info("zone2_prepared", project_id=project.id, release_id=release_id, **counts)
    return release


def _name_day(value) -> str:
    text = str(value or "")
    if len(text) == 8 and text.isdigit():
        return f"{text[:4]}-{text[4:6]}-{text[6:]}"
    return ""


def _release_redcap(cursor, cohort, archive: Path, safe: Path) -> List[List[str]]:
    members = {m.patient_id: m for m in cohort}
    codes = list(members)
    rows_out: List[List[str]] = []
    for dataset in redcap.datasets_with_values(cursor, codes):
        defined = {f.field_name: f for f in redcap.fields(cursor, dataset)}
        columns, rows = redcap.wide_rows(cursor, dataset, codes)
        _write_csv(
            archive / "original" / f"redcap_{dataset}.csv",
            ["patient_id", "event", "instance"] + columns,
            rows,
        )
        kept = []
        for column in columns:
            field = defined.get(column.split("___")[0])
            kind = redcap.release_kind(field) if field else None
            if kind:
                kept.append((column, kind, field))
        name = f"redcap_{dataset}.csv"
        released = []
        for row in rows:
            member = members[row["patient_id"]]
            out = {
                "project_patient_id": member.project_patient_id,
                "event": row.get("event", ""),
                "instance": row.get("instance", ""),
            }
            for column, kind, _ in kept:
                value = row.get(column, "")
                out[column] = shift_date(value, member.date_shift_days) if kind == "date" else value
            released.append(out)
        _write_csv(
            safe / name,
            ["project_patient_id", "event", "instance"] + [c for c, _, _ in kept],
            released,
        )
        rows_out += [
            [name, "project_patient_id", "text", "The patient's ID in this project only", "See patients.csv"],
            [name, "event", "text", "REDCap event, if the project has events", ""],
            [name, "instance", "text", "Repeat number, if the form repeats", ""],
        ]
        for column, kind, field in kept:
            kinds = field.field_type + (f" ({field.validation})" if field.validation else "")
            notes = data_dictionary.SHIFTED if kind == "date" else (field.choices or "")
            rows_out.append([name, column, kinds, field.label or "", notes])
    return rows_out


def _build(cursor, project, cohort, release_id, archive: Path, safe: Path) -> dict:
    original = {"patients": [], "applications": [], "documents": []}
    zone2 = {"patients": [], "applications": [], "documents": []}
    application_map, document_map = [], []

    for member in cohort:
        patient = patients_crud.get_patient(cursor, member.patient_id)
        if patient is None:
            continue
        days, pid = member.date_shift_days, member.project_patient_id
        original["patients"].append(patient.model_dump(mode="json"))
        zone2["patients"].append(
            {
                "project_patient_id": pid,
                "age_at_registration": age_at(patient.dt_b, patient.dt_reg),
                "registration_date": shift_date(patient.dt_reg, days),
                "death_date": shift_date(patient.dt_d, days),
                "state": patient.ptstate,
                "country": patient.ptcountry,
            }
        )

        for application in _submitted_applications(cursor, member.patient_id):
            app_id = f"{project.short_code}-A{projects_crud.random_code(6)}"
            application_map.append({"application_id": application.id, "project_application_id": app_id})
            original["applications"].append(application.model_dump(mode="json"))
            zone2["applications"].append(
                {
                    "project_application_id": app_id,
                    "project_patient_id": pid,
                    "status": application.status,
                    "created_date": shift_date(application.created_at, days),
                    "submitted_date": shift_date(application.submitted_at, days),
                }
            )

            for record in _documents(cursor, application.id):
                doc_id = f"{project.short_code}-D{projects_crud.random_code(6)}"
                document_map.append({"file_id": record.id, "project_document_id": doc_id})
                original["documents"].append(record.model_dump(mode="json"))
                released_file = ""
                if project.documents_approved:
                    try:
                        name = _copy_document(
                            record, safe / "documents" / pid, doc_id,
                            member.patient_id, pid, days, uid_salt=project.id,
                        )
                    except DocumentError as exc:
                        raise DocumentError(f"{record.original_file_name}: {exc}") from exc
                    except Exception as exc:
                        raise DocumentError(f"{record.original_file_name}: could not be prepared ({exc})") from exc
                    released_file = f"documents/{pid}/{name}"
                zone2["documents"].append(
                    {
                        "project_document_id": doc_id,
                        "project_application_id": app_id,
                        "project_patient_id": pid,
                        "file_type": (record.file_extension or "").lower(),
                        "modality": record.modality or "",
                        "document_date": shift_date(_name_day(record.document_date), days),
                        "file_size_bytes": record.file_size,
                        "created_date": shift_date(record.created_at, days),
                        "file": released_file,
                    }
                )

    for name, rows in original.items():
        columns = sorted({k for row in rows for k in row}) or ["empty"]
        _write_csv(archive / "original" / f"{name}.csv", columns, rows)
    _write_csv(
        archive / "id_map.csv",
        ["patient_id", "project_patient_id", "date_shift_days"],
        _model_rows(cohort),
    )
    _write_csv(archive / "application_map.csv", ["application_id", "project_application_id"], application_map)
    _write_csv(archive / "document_map.csv", ["file_id", "project_document_id"], document_map)

    _write_csv(safe / "patients.csv", PATIENT_COLUMNS, zone2["patients"])
    _write_csv(safe / "applications.csv", APPLICATION_COLUMNS, zone2["applications"])
    _write_csv(safe / "documents.csv", DOCUMENT_COLUMNS, zone2["documents"])
    dictionary_rows = data_dictionary.release_rows()
    dictionary_rows += _release_redcap(cursor, cohort, archive, safe)
    _write_csv(
        safe / "data_dictionary.csv",
        data_dictionary.RELEASE_HEADER,
        [dict(zip(data_dictionary.RELEASE_HEADER, row)) for row in dictionary_rows],
    )
    (safe / "README.txt").write_text(
        README.format(
            name=project.name,
            code=project.short_code,
            release=release_id,
            cap=AGE_CAP,
            documents_note=(
                "\ndocuments/        the de-identified files, by project patient ID"
                if project.documents_approved
                else "\n                  (the files themselves are not included)"
            ),
        )
    )

    return {
        "patients": len(zone2["patients"]),
        "applications": len(zone2["applications"]),
        "documents": len(zone2["documents"]),
    }


def summary(release: ProjectRelease) -> Dict[str, dict]:
    folder = Path(release.release_path or release.inbox_path or "")
    columns, samples = {}, {}
    for name in ("patients", "applications", "documents"):
        path = folder / f"{name}.csv"
        if not path.is_file():
            continue
        header = path.read_text(encoding="utf-8").splitlines()[:1]
        columns[name] = next(csv.reader(io.StringIO(header[0]))) if header else []
        samples[name] = _read_csv(path)[:SAMPLE_ROWS]
    return {"columns": columns, "samples": samples}


def _remove_from_inbox(folder: Path) -> None:
    if not str(folder) or not folder.resolve().is_relative_to(storage.zone2_inbox_root().resolve()):
        return
    shutil.rmtree(folder, ignore_errors=True)
    try:
        folder.parent.rmdir()
    except OSError:
        pass


def _pending(cursor, release_id: str) -> ProjectRelease:
    release = projects_crud.get_release_or_404(cursor, release_id)
    if release.status != "in_review":
        raise ValidationError(f"This release is already '{release.status}'")
    return release


def _refuse_leaded_since_prepared(cursor, release: ProjectRelease) -> None:
    if not release.files_included:
        return
    path = Path(release.archive_path or "") / "document_map.csv"
    released = {row["file_id"] for row in _read_csv(path)} if path.is_file() else set()
    leaded = [
        record.deidentified_file_name or record.original_file_name
        for record in files_crud.list_files(cursor)
        if record.id in released and record.leaded_phi
    ]
    if leaded:
        raise ValidationError(
            f"{len(leaded)} image(s) in this release were marked as having leaded "
            "identifiers after it was prepared; reject it and prepare again: "
            + ", ".join(leaded[:10])
        )


def approve(cursor, release_id: str, actor_id: str) -> ProjectRelease:
    release = _pending(cursor, release_id)
    project = projects_crud.get_project_or_404(cursor, release.project_id)
    inbox = Path(release.inbox_path or "")
    if not inbox.is_dir():
        raise ValidationError("The safe copy is no longer in the inbox; prepare the project again")
    _refuse_leaded_since_prepared(cursor, release)

    target = storage.zone2_root() / project.short_code / release.id
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copytree(inbox, target)
    _remove_from_inbox(inbox)

    release = projects_crud.decide_release(cursor, release.id, "approved", actor_id, release_path=str(target))
    projects_crud.set_project_status(cursor, project.id, "released")
    log.info("zone2_released", project_id=project.id, release_id=release.id)
    return release


def reject(cursor, release_id: str, actor_id: str, reason: Optional[str]) -> ProjectRelease:
    text = (reason or "").strip()
    if not text:
        raise ValidationError("A reason is required to reject a release")
    release = _pending(cursor, release_id)
    _remove_from_inbox(Path(release.inbox_path or ""))

    release = projects_crud.decide_release(cursor, release.id, "rejected", actor_id, reason=text)
    projects_crud.set_project_status(cursor, release.project_id, "rejected", text)
    log.info("zone2_rejected", project_id=release.project_id, release_id=release.id)
    return release


def _release_folder(release: ProjectRelease) -> Path:
    return Path(release.release_path or release.inbox_path or "")


def release_documents(release: ProjectRelease) -> List[dict]:
    path = _release_folder(release) / "documents.csv"
    if not path.is_file():
        return []
    return [row for row in _read_csv(path) if row.get("file")]


def release_file(release: ProjectRelease, document_id: str) -> Path:
    folder = _release_folder(release).resolve()
    for row in release_documents(release):
        if row.get("project_document_id") == document_id:
            path = (folder / row["file"]).resolve()
            if path.is_relative_to(folder) and path.is_file():
                return path
            break
    raise NotFoundError(f"Document '{document_id}' is not in this release")


def _source_file_id(release: ProjectRelease, document_id: str) -> str:
    path = Path(release.archive_path or "") / "document_map.csv"
    for row in _read_csv(path) if path.is_file() else []:
        if row.get("project_document_id") == document_id:
            return row["file_id"]
    raise NotFoundError(f"Document '{document_id}' is not in this release's archive")


def flag_document(cursor, release_id: str, document_id: str, actor_id: str, note: Optional[str]):
    text = (note or "").strip()
    if not text:
        raise ValidationError("Say what is wrong with the document")
    release = _pending(cursor, release_id)
    release_file(release, document_id)
    file_id = _source_file_id(release, document_id)

    record = files_crud.get_file_or_404(cursor, file_id)
    files_crud.update_file(
        cursor,
        file_id,
        PatientApplicationFileUpdate(
            review_status="rejected",
            review_note=f"Flagged in Zone 2 release {release.id} ({document_id}): {text}",
        ),
    )
    release = reject(cursor, release.id, actor_id, f"Document {document_id} flagged: {text}")
    log.info("zone2_document_flagged", release_id=release.id, document_id=document_id, file_id=file_id)
    return release, record


def end_project(cursor, project_id: str, reason: Optional[str]) -> Project:
    text = (reason or "").strip()
    if not text:
        raise ValidationError("A reason is required to end a project")
    project = projects_crud.get_project_or_404(cursor, project_id)
    if project.status == ENDED:
        raise ValidationError("The project has already ended")
    if project.status == "in_review":
        raise ValidationError("A release is waiting for review; approve or reject it first")
    project = projects_crud.set_project_status(cursor, project.id, ENDED, text)
    log.info("project_ended", project_id=project.id)
    return project
