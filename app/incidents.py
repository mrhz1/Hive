import io
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any, List, Optional

from pydantic import BaseModel, Field, computed_field, field_validator

from app import lockdown
from app.db import NOW_SQL, execute
from app.errors import NotFoundError, ValidationError
from app.logging_setup import get_logger
from app.mailer import send_email

log = get_logger(__name__)

PRELIMINARY_HOURS = 24
FULL_HOURS = 72
REMIND_BEFORE_HOURS = 6

OPEN, CLOSED = "open", "closed"
TEXT_FIELDS = ("description", "data_affected", "risk", "containment", "affected_parties")

COLUMNS = (
    "id", "title", "level", "status", "detected_at", "description", "data_affected",
    "risk", "containment", "affected_parties", "preliminary_sent_at", "full_sent_at",
    "preliminary_reminded_at", "full_reminded_at", "created_by_id", "created_at",
    "updated_at", "closed_at",
)
_COLS = ", ".join(f"`{c}`" for c in COLUMNS)


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _iso(moment: datetime) -> str:
    return moment.isoformat(timespec="seconds")


def _parse(value: Optional[str]) -> Optional[datetime]:
    if not value:
        return None
    try:
        moment = datetime.fromisoformat(value)
    except ValueError:
        return None
    return moment if moment.tzinfo else moment.replace(tzinfo=timezone.utc)


class IncidentFields(BaseModel):
    title: Optional[str] = Field(default=None, max_length=200)
    level: Optional[str] = Field(default=None, pattern="^(low|high|critical)$")
    detected_at: Optional[datetime] = None
    description: Optional[str] = None
    data_affected: Optional[str] = None
    risk: Optional[str] = None
    containment: Optional[str] = None
    affected_parties: Optional[str] = None

    @field_validator("detected_at")
    @classmethod
    def _not_future(cls, value: Optional[datetime]) -> Optional[datetime]:
        if value is None:
            return None
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        if value > _now() + timedelta(minutes=5):
            raise ValueError("The detection time is in the future")
        return value


class IncidentCreate(IncidentFields):
    title: str = Field(min_length=1, max_length=200)
    level: str = Field(pattern="^(low|high|critical)$")
    detected_at: datetime


class Incident(BaseModel):
    id: str
    title: str
    level: str
    status: str
    detected_at: str
    description: Optional[str] = None
    data_affected: Optional[str] = None
    risk: Optional[str] = None
    containment: Optional[str] = None
    affected_parties: Optional[str] = None
    preliminary_sent_at: Optional[str] = None
    full_sent_at: Optional[str] = None
    created_by_id: Optional[str] = None
    created_at: Optional[Any] = None
    updated_at: Optional[str] = None
    closed_at: Optional[str] = None

    def _due(self, hours: int) -> Optional[datetime]:
        detected = _parse(self.detected_at)
        return detected + timedelta(hours=hours) if detected else None

    @computed_field
    @property
    def preliminary_due(self) -> Optional[str]:
        due = self._due(PRELIMINARY_HOURS)
        return _iso(due) if due else None

    @computed_field
    @property
    def full_due(self) -> Optional[str]:
        due = self._due(FULL_HOURS)
        return _iso(due) if due else None


def _row(row) -> Incident:
    values = dict(zip(COLUMNS, row))
    values.pop("preliminary_reminded_at", None)
    values.pop("full_reminded_at", None)
    return Incident(**values)


def _reminded(cursor, incident_id: str) -> dict:
    execute(
        cursor,
        "SELECT `preliminary_reminded_at`, `full_reminded_at` FROM `incidents` WHERE `id` = %s",
        (incident_id,),
    )
    row = cursor.fetchone() or (None, None)
    return {"preliminary": row[0], "full": row[1]}


def list_incidents(cursor) -> List[Incident]:
    execute(cursor, f"SELECT {_COLS} FROM `incidents` ORDER BY `created_at` DESC")
    return [_row(r) for r in cursor.fetchall()]


def get_or_404(cursor, incident_id: str) -> Incident:
    execute(cursor, f"SELECT {_COLS} FROM `incidents` WHERE `id` = %s", (incident_id,))
    row = cursor.fetchone()
    if not row:
        raise NotFoundError(f"Incident '{incident_id}' not found")
    return _row(row)


def _update(cursor, incident_id: str, fields: dict) -> Incident:
    fields = dict(fields, updated_at=_iso(_now()))
    parts = ", ".join(f"`{c}` = %s" for c in fields)
    execute(
        cursor,
        f"UPDATE `incidents` SET {parts} WHERE `id` = %s",
        tuple(fields.values()) + (incident_id,),
    )
    return get_or_404(cursor, incident_id)


def create(cursor, payload: IncidentCreate, actor_id: str) -> Incident:
    incident_id = str(uuid.uuid4())
    values = {c: None for c in COLUMNS}
    values.update(
        payload.model_dump(exclude={"detected_at"}),
        id=incident_id,
        status=OPEN,
        detected_at=_iso(payload.detected_at),
        created_by_id=actor_id,
        updated_at=_iso(_now()),
    )
    slots = ", ".join(NOW_SQL if c == "created_at" else "%s" for c in COLUMNS)
    execute(
        cursor,
        f"INSERT INTO `incidents` ({_COLS}) VALUES ({slots})",
        tuple(values[c] for c in COLUMNS if c != "created_at"),
    )
    incident = get_or_404(cursor, incident_id)
    log.warning("incident_recorded", incident_id=incident_id, level=incident.level)
    if incident.level in ("high", "critical"):
        _announce(incident)
    return incident


def update(cursor, incident_id: str, payload: IncidentFields) -> Incident:
    incident = get_or_404(cursor, incident_id)
    if incident.status == CLOSED:
        raise ValidationError("The incident is closed")
    fields = payload.model_dump(exclude_unset=True)
    if "detected_at" in fields and fields["detected_at"] is not None:
        fields["detected_at"] = _iso(fields["detected_at"])
    fields = {k: v for k, v in fields.items() if v is not None or k in TEXT_FIELDS}
    return _update(cursor, incident_id, fields) if fields else incident


def mark_sent(cursor, incident_id: str, report: str) -> Incident:
    incident = get_or_404(cursor, incident_id)
    column = {"preliminary": "preliminary_sent_at", "full": "full_sent_at"}.get(report)
    if column is None:
        raise ValidationError("report must be 'preliminary' or 'full'")
    if getattr(incident, column):
        raise ValidationError(f"The {report} report is already marked as sent")
    return _update(cursor, incident_id, {column: _iso(_now())})


def close(cursor, incident_id: str) -> Incident:
    incident = get_or_404(cursor, incident_id)
    if incident.status == CLOSED:
        raise ValidationError("The incident is already closed")
    missing = [
        label for label, name in (("containment", "containment"), ("data affected", "data_affected"))
        if not (getattr(incident, name) or "").strip()
    ]
    if missing:
        raise ValidationError(f"Fill in {', '.join(missing)} before closing")
    if not incident.full_sent_at:
        raise ValidationError("Mark the full report as sent before closing")
    return _update(cursor, incident_id, {"status": CLOSED, "closed_at": _iso(_now())})


def _announce(incident: Incident) -> None:
    send_email(
        lockdown.contacts(incident.level),
        f"Security incident ({incident.level}): {incident.title}",
        "\n".join([
            f"A {incident.level} security incident was recorded in Hive.",
            "",
            f"Title: {incident.title}",
            f"Detected: {incident.detected_at}",
            f"Preliminary report due: {incident.preliminary_due}",
            f"Full report due: {incident.full_due}",
            "",
            incident.description or "",
            "",
            "Consider locking the app from the Security page if it is still going on.",
            "",
            "-- Hive",
        ]),
    )


def due_reminders(cursor, now: Optional[datetime] = None) -> List[str]:
    now = now or _now()
    sent = []
    for incident in list_incidents(cursor):
        if incident.status != OPEN:
            continue
        reminded = _reminded(cursor, incident.id)
        for report, hours, done in (
            ("preliminary", PRELIMINARY_HOURS, incident.preliminary_sent_at),
            ("full", FULL_HOURS, incident.full_sent_at),
        ):
            due = incident._due(hours)
            if done or reminded[report] or due is None:
                continue
            if now < due - timedelta(hours=REMIND_BEFORE_HOURS):
                continue
            late = now >= due
            send_email(
                lockdown.contacts("high"),
                f"{'OVERDUE' if late else 'Due soon'}: {report} breach report - {incident.title}",
                "\n".join([
                    f"The {report} report for the incident '{incident.title}' "
                    f"{'was due' if late else 'is due'} at {_iso(due)}.",
                    "",
                    "Send it, then mark it as sent on the Security page.",
                    "",
                    "-- Hive",
                ]),
            )
            execute(
                cursor,
                f"UPDATE `incidents` SET `{report}_reminded_at` = %s WHERE `id` = %s",
                (_iso(now), incident.id),
            )
            sent.append(f"{incident.id}:{report}")
    return sent


def report_docx(incident: Incident) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(f"Security incident report: {incident.title}", level=1)
    table = document.add_table(rows=0, cols=2)
    for label, value in (
        ("Level", incident.level),
        ("Status", incident.status),
        ("Detected", incident.detected_at),
        ("Preliminary report due", incident.preliminary_due),
        ("Preliminary report sent", incident.preliminary_sent_at or "not yet"),
        ("Full report due", incident.full_due),
        ("Full report sent", incident.full_sent_at or "not yet"),
    ):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value or ""
    for label, value in (
        ("What happened", incident.description),
        ("Data affected (scope)", incident.data_affected),
        ("Risk assessment", incident.risk),
        ("Containment actions", incident.containment),
        ("Affected parties", incident.affected_parties),
    ):
        document.add_heading(label, level=2)
        document.add_paragraph(value or "Not recorded yet.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
