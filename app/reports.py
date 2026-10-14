"""Writing the big operational reports (migration, QC): a folder holding an
Excel workbook (a summary sheet, then one sheet per list) and every list in
full as CSV.

Excel holds about a million rows, and a migration can have millions of
problem files, so each list sheet stops at `cap` rows and says so; the CSV
next to it always has them all. The workbook is written in streaming mode,
so memory stays flat.
"""

import csv
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from app.safe_cells import row as safe_row
from app.storage import REPO_ROOT

DEFAULT_CAP = 50_000


def reports_root() -> Path:
    value = Path(os.environ.get("REPORTS_DIR", "storage/reports"))
    return value if value.is_absolute() else REPO_ROOT / value


def new_folder(kind: str, root: Path = None) -> Path:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    folder = (root or reports_root()) / f"{kind}-{stamp}"
    folder.mkdir(parents=True, exist_ok=False)
    return folder


class ListOut:
    """One list: written to CSV in full as rows come, and kept (up to the
    cap) for the workbook."""

    def __init__(self, folder: Path, name: str, title: str, columns: Sequence[str], cap: int):
        self.name, self.title, self.columns, self.cap = name, title, list(columns), cap
        self.count = 0
        self.kept: List[list] = []
        self._handle = open(folder / f"{name}.csv", "w", newline="", encoding="utf-8")
        self._writer = csv.writer(self._handle)
        self._writer.writerow(self.columns)

    def add(self, row: Sequence) -> None:
        values = safe_row(["" if v is None else v for v in row])
        self._writer.writerow(values)
        self.count += 1
        if len(self.kept) < self.cap:
            self.kept.append(values)

    def close(self) -> None:
        self._handle.close()


def write_workbook(
    path: Path,
    title: str,
    summary: Iterable[Tuple[str, str, object]],
    lists: Sequence[ListOut],
) -> Path:
    """Summary rows are (section, measure, value)."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    book = Workbook(write_only=True)
    bold = Font(bold=True)

    def header(sheet, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            cell.font = bold
            cells.append(cell)
        sheet.append(cells)

    sheet = book.create_sheet("Summary")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 60
    sheet.column_dimensions["C"].width = 16
    sheet.append([title])
    sheet.append([f"Made {datetime.now(timezone.utc).isoformat(timespec='seconds')}"])
    sheet.append([])
    header(sheet, ["Section", "Measure", "Count"])
    for row in summary:
        sheet.append(safe_row(row))
    sheet.append([])
    header(sheet, ["List", "Rows", "File with every row"])
    for out in lists:
        sheet.append([out.title, out.count, f"{out.name}.csv"])

    for out in lists:
        tab = book.create_sheet("".join(c for c in out.title if c not in r"[]:*?/\\")[:31])
        if out.count > len(out.kept):
            tab.append([f"First {len(out.kept)} of {out.count} rows; all are in {out.name}.csv"])
        header(tab, out.columns)
        for row in out.kept:
            tab.append(row)

    book.save(str(path))
    return path
