"""Loading patients in bulk from a spreadsheet (Excel .xlsx or CSV).

The first row names the columns: the template's labels ("First name",
"Date of birth", ...) or the database names (fstname, dt_b, ...). Only
"Patient code" is required.

- a code not in the system: the patient is created;
- a code already there: only its empty fields are filled; a value that
  differs from what is stored is reported as a conflict and left alone;
- rows with problems (bad code, bad date, bad email, a code twice, an email
  or phone of another patient) are skipped and reported.

Dates are real Excel dates or YYYY-MM-DD; 03/12/1980 could be either day or
month first, so it is refused. Imported patients' details count as not
checked: their first application waits until someone confirms them.
"""

import csv
import io
import re
from itertools import islice
from datetime import date, datetime
from typing import Dict, List, Optional, Tuple

from pydantic import BaseModel, Field
from pydantic import ValidationError as PydanticError

from app.crud import patients as crud
from app.errors import ValidationError
from app.ids import clean_patient_code, is_patient_code
from app.logging_setup import get_logger
from app.schemas import PatientUpdate, patient_code_invalid

log = get_logger(__name__)

MAX_ROWS = 20000

# (column, label) in template order. Patient fields first, then the
# provider's (the clinic/physician).
FIELDS: List[Tuple[str, str]] = [
    ("id", "Patient code"),
    ("fstname", "First name"),
    ("lstname", "Last name"),
    ("dt_b", "Date of birth"),
    ("dt_d", "Date of death"),
    ("dt_reg", "Registration date"),
    ("ptemail", "Email"),
    ("ptphone", "Phone"),
    ("ptphone2", "Phone 2"),
    ("ptwphone", "Work phone"),
    ("ptwphone2", "Work phone 2"),
    ("ptstreet", "Street"),
    ("ptstreet2", "Street 2"),
    ("ptstreet3", "Street 3"),
    ("ptcity", "City"),
    ("ptstate", "State"),
    ("ptzip", "ZIP"),
    ("ptcountry", "Country"),
    ("instcode", "Institution code"),
    ("pname", "Provider name"),
    ("pemail", "Provider email"),
    ("phone1", "Provider phone 1"),
    ("phone2", "Provider phone 2"),
    ("wphone1", "Provider work phone 1"),
    ("wphone2", "Provider work phone 2"),
    ("street", "Provider street"),
    ("street2", "Provider street 2"),
    ("street3", "Provider street 3"),
    ("city", "Provider city"),
    ("state", "Provider state"),
    ("zip", "Provider ZIP"),
    ("country", "Provider country"),
]
DATE_FIELDS = {"dt_b", "dt_d", "dt_reg"}
UNIQUE_FIELDS = (("ptemail", "Email"), ("ptphone", "Phone"))

CREATE, UPDATE, UNCHANGED, ERROR = "create", "update", "unchanged", "error"


class ImportRow(BaseModel):
    line: int
    code: str = ""
    action: str
    filled: List[str] = Field(default_factory=list)
    conflicts: List[str] = Field(default_factory=list)
    errors: List[str] = Field(default_factory=list)


class ImportResult(BaseModel):
    applied: bool = False
    counts: Dict[str, int]
    ignored_columns: List[str] = Field(default_factory=list)
    rows: List[ImportRow]


def _key(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


# Labels win over column names: "City" is the patient's (ptcity), though the
# provider's column is called city.
_HEADERS = {_key(column): column for column, _ in FIELDS}
_HEADERS.update({_key(label): column for column, label in FIELDS})
_HEADERS.update({_key("code"): "id", _key("patient id"): "id", _key("dob"): "dt_b"})

_LABELS = dict(FIELDS)


def template_csv() -> bytes:
    return ("﻿" + ",".join(label for _, label in FIELDS) + "\r\n").encode("utf-8")


# --- reading the file -------------------------------------------------------


def _read_rows(name: str, data: bytes, max_rows: Optional[int] = None) -> List[List[object]]:
    """The rows of the first sheet (or the CSV). With `max_rows`, stops one
    row past it, so a file far too big is refused without being read whole
    into memory; the caller still sees there are too many."""
    limit = None if max_rows is None else max_rows + 2  # the header, one too many
    if (name or "").lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise ValidationError(f"Could not read the Excel file: {e}")
        sheet = book.worksheets[0]
        return [list(row) for row in islice(sheet.iter_rows(values_only=True), limit)]

    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in islice(csv.reader(io.StringIO(text), dialect), limit)]


def _text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _date(value) -> Tuple[Optional[date], Optional[str]]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    text = str(value).strip()
    try:
        return date.fromisoformat(text[:10]), None
    except ValueError:
        return None, f"'{text}' is not a date (use YYYY-MM-DD)"


def _parse(row: List[object], columns: List[Optional[str]]) -> Tuple[Dict, List[str]]:
    values: Dict = {}
    errors: List[str] = []
    for column, raw in zip(columns, row):
        if column is None:
            continue
        if column in DATE_FIELDS:
            day, error = _date(raw)
            if error:
                errors.append(f"{_LABELS[column]}: {error}")
            elif day:
                values[column] = day
        else:
            text = _text(raw)
            if text is not None:
                values[column] = text

    code = clean_patient_code(values.pop("id", None) or "")
    if not code:
        errors.append("Patient code is missing")
    elif not is_patient_code(code):
        errors.append(patient_code_invalid(code))

    try:
        checked = PatientUpdate(**values).model_dump(exclude_unset=True)
        values = {k: v for k, v in checked.items() if v is not None}
    except PydanticError as e:
        for problem in e.errors():
            field = str(problem["loc"][0]) if problem.get("loc") else ""
            errors.append(f"{_LABELS.get(field, field)}: {problem['msg']}")
    return dict(values, id=code), errors


# --- planning ---------------------------------------------------------------------


def _same(a, b) -> bool:
    return str(a).strip().lower() == str(b).strip().lower()


def plan(cursor, name: str, data: bytes):
    """What importing would do, row by row: (result, new rows, fills)."""
    rows = _read_rows(name, data, MAX_ROWS)
    while rows and not any(_text(v) for v in rows[0]):
        rows.pop(0)
    if not rows:
        raise ValidationError("The file is empty")

    header = rows[0]
    columns = [_HEADERS.get(_key(h)) for h in header]
    if "id" not in columns:
        raise ValidationError("No 'Patient code' column in the first row")
    ignored = [str(h) for h, c in zip(header, columns) if c is None and _text(h)]

    body = [(n, r) for n, r in enumerate(rows[1:], start=2) if any(_text(v) for v in r)]
    if len(body) > MAX_ROWS:
        raise ValidationError(
            f"The file has more than {MAX_ROWS} rows; split it into smaller "
            "files and import them one by one"
        )

    existing = {p.id: p for p in crud.list_patients(cursor)}
    owners: Dict[Tuple[str, str], str] = {}
    for patient in existing.values():
        for field, _ in UNIQUE_FIELDS:
            value = getattr(patient, field, None)
            if value:
                owners[(field, str(value).strip().lower())] = patient.id

    seen: Dict[str, int] = {}
    result_rows: List[ImportRow] = []
    new_rows: List[Dict] = []
    fills: Dict[str, Dict] = {}

    for line, raw in body:
        values, errors = _parse(raw, columns)
        code = values["id"]

        if code and code in seen:
            errors.append(f"Patient code {code} is also on line {seen[code]}")
        elif code:
            seen[code] = line

        for field, label in UNIQUE_FIELDS:
            value = values.get(field)
            owner = owners.get((field, str(value).strip().lower())) if value else None
            if owner and owner != code:
                errors.append(f"{label} '{value}' belongs to patient {owner}")

        if errors:
            result_rows.append(ImportRow(line=line, code=code, action=ERROR, errors=errors))
            continue

        for field, _ in UNIQUE_FIELDS:
            if values.get(field):
                owners[(field, str(values[field]).strip().lower())] = code

        current = existing.get(code)
        if current is None:
            new_rows.append(values)
            result_rows.append(
                ImportRow(
                    line=line, code=code, action=CREATE,
                    filled=[_LABELS[k] for k in values if k != "id"],
                )
            )
            continue

        fill, conflicts = {}, []
        for field, value in values.items():
            if field == "id":
                continue
            stored = getattr(current, field, None)
            if stored in (None, ""):
                fill[field] = value
            elif not _same(stored, value):
                conflicts.append(f"{_LABELS[field]}: stored '{stored}', file '{value}'")
        if fill:
            fills[code] = fill
        result_rows.append(
            ImportRow(
                line=line, code=code, action=UPDATE if fill else UNCHANGED,
                filled=[_LABELS[k] for k in fill], conflicts=conflicts,
            )
        )

    counts = {action: 0 for action in (CREATE, UPDATE, UNCHANGED, ERROR)}
    for row in result_rows:
        counts[row.action] += 1
    result = ImportResult(counts=counts, ignored_columns=ignored, rows=result_rows)
    return result, new_rows, fills


def run(cursor, name: str, data: bytes) -> ImportResult:
    """Create the new patients and fill the existing ones' empty fields;
    rows with errors are skipped."""
    result, new_rows, fills = plan(cursor, name, data)
    if new_rows:
        crud.insert_many(cursor, new_rows)
    for code, fill in fills.items():
        crud.set_fields(cursor, code, fill)
    result.applied = True
    log.info("patients_imported", **result.counts)
    return result
